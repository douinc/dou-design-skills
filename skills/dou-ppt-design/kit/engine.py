"""
슬라이드 빌드 엔진 — 같은 레이아웃 정의에서 (1) 파워포인트(.pptx) 와 (2) 미리보기 HTML 을 함께 만든다.

좌표계: 1280 × 720 px (16:9, 96 px/in → 13.333 × 7.5 in). pptx 에는 EMU 로 환산(1 px = 9525 EMU).
글꼴: Noto Sans KR (구글 슬라이드에 기본 탑재). 미리보기는 Google Fonts 에서 같은 글꼴을 불러온다.
"""
from __future__ import annotations

import html as _html
import os
import re
from dataclasses import dataclass, field

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

PX = 9525  # EMU per px @96dpi
W, H = 1280, 720
# 구글 슬라이드 기본 탑재 글꼴. 파워포인트로만 볼 사람이 있으면
#   DECK_FONT="맑은 고딕" python3 build.py  처럼 바꿔서 한 벌 더 뽑는다.
FONT = os.environ.get("DECK_FONT", "Noto Sans KR")
USED_CHARS: set = set()

# 발표장에서 읽히도록 본문 글자를 일괄로 키운다. 제목(22px 이상)은 그대로.
TEXT_BUMP = float(os.environ.get("TEXT_BUMP", "1.2"))


def fs(px: float) -> float:
    return px if px >= 22 else round(px + TEXT_BUMP, 1)
ASSETS = os.path.join(os.path.dirname(os.path.abspath(__file__)), "assets")


def hexrgb(h: str) -> RGBColor:
    h = h.lstrip("#")
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


# ---------------------------------------------------------------- primitives
@dataclass
class Run:
    text: str
    bold: bool = False
    color: str | None = None
    size: float | None = None
    weight: int | None = None  # 400/500/700 — pptx 는 bold 여부만 반영

    def __post_init__(self):
        if self.size is not None:
            self.size = fs(self.size)


@dataclass
class Text:
    x: float
    y: float
    w: float
    h: float
    runs: list  # list[Run] | list[list[Run]] (paragraphs)
    size: float = 14
    color: str = "#1A2233"
    weight: int = 400
    align: str = "left"  # left|center|right
    valign: str = "top"  # top|middle|bottom
    line: float = 1.35
    para_gap: float = 0  # px between paragraphs
    letter: float = 0  # letter-spacing px (html only)
    name: str = ""

    def __post_init__(self):
        self.size = fs(self.size)


@dataclass
class Rect:
    x: float
    y: float
    w: float
    h: float
    fill: str | None = "#FFFFFF"
    radius: float = 0
    stroke: str | None = None
    stroke_w: float = 1
    dash: bool = False
    shadow: bool = False
    shape: str = "rect"  # rect|ellipse
    opacity: float = 1.0
    name: str = ""


@dataclass
class Line:
    x1: float
    y1: float
    x2: float
    y2: float
    color: str = "#E3E8F0"
    w: float = 1
    dash: bool = False
    arrow: bool = False


@dataclass
class Image:
    path: str  # png 경로 (assets 상대)
    x: float
    y: float
    w: float
    h: float
    name: str = ""


@dataclass
class Icon:
    """assets/icons/<name>-<color>.png 로 렌더된 아이콘. html 에서는 svg 인라인."""
    name: str
    x: float
    y: float
    size: float = 24
    color: str = "#3B82F6"


@dataclass
class Slide:
    title: str
    items: list = field(default_factory=list)
    bg: str = "#FFFFFF"
    notes: str = ""
    section: str = ""
    balance: bool = True  # 아래 여백을 본문 위 간격으로 옮길지

    def add(self, *items):
        for it in items:
            if isinstance(it, (list, tuple)):
                self.add(*it)
            elif it is not None:
                self.items.append(it)
        return self


# ------------------------------------------------------------------- helpers
def runs(*parts) -> list[Run]:
    """runs("일반 ", ("굵게", {"bold": True}), " 이어서") 형식."""
    out = []
    for p in parts:
        if isinstance(p, Run):
            out.append(p)
        elif isinstance(p, tuple):
            out.append(Run(p[0], **p[1]))
        else:
            out.append(Run(str(p)))
    return out


def paras(*lines) -> list[list[Run]]:
    return [runs(*l) if isinstance(l, (list, tuple)) and not (len(l) == 2 and isinstance(l[1], dict)) else runs(l) for l in lines]


def _paragraphs(t: Text) -> list[list[Run]]:
    if not t.runs:
        return [[]]
    if isinstance(t.runs, str):
        paras_ = [[Run(t.runs)]]
    elif all(isinstance(r, Run) for r in t.runs):
        paras_ = [t.runs]
    else:
        paras_ = [r if isinstance(r, list) else [r] for r in t.runs]
    # 글 안의 "\n" 은 줄바꿈(새 문단)으로 처리한다
    out = []
    for para in paras_:
        cur = []
        for r in para:
            pieces = r.text.split("\n")
            for k, piece in enumerate(pieces):
                if k:
                    out.append(cur)
                    cur = []
                cur.append(Run(piece, r.bold, r.color, None, r.weight) if len(pieces) > 1 else r)
                if len(pieces) > 1:
                    cur[-1].size = r.size
        out.append(cur)
    return out


# ---------------------------------------------------------------- pptx emit
class PptxWriter:
    def __init__(self):
        self.prs = Presentation()
        self.prs.slide_width = Emu(W * PX)
        self.prs.slide_height = Emu(H * PX)
        self.blank = self.prs.slide_layouts[6]

    def _rgb_fill(self, shape, color, opacity=1.0):
        if color is None:
            shape.fill.background()
            return
        shape.fill.solid()
        shape.fill.fore_color.rgb = hexrgb(color)
        if opacity < 1.0:
            srgb = shape.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            a = etree.SubElement(srgb, qn("a:alpha"))
            a.set("val", str(int(opacity * 100000)))

    def _line(self, shape, color, w=1.0, dash=False):
        if color is None:
            shape.line.fill.background()
            return
        shape.line.color.rgb = hexrgb(color)
        shape.line.width = Emu(int(w * PX))
        if dash:
            ln = shape.line._get_or_add_ln()
            pd = etree.SubElement(ln, qn("a:prstDash"))
            pd.set("val", "dash")

    def _no_shadow(self, shape):
        spPr = shape._element.spPr
        eff = etree.SubElement(spPr, qn("a:effectLst"))
        return eff

    def _shadow(self, shape):
        spPr = shape._element.spPr
        eff = etree.SubElement(spPr, qn("a:effectLst"))
        sh = etree.SubElement(eff, qn("a:outerShdw"))
        sh.set("blurRad", str(14 * PX))
        sh.set("dist", str(4 * PX))
        sh.set("dir", "5400000")
        sh.set("algn", "t")
        sh.set("rotWithShape", "0")
        c = etree.SubElement(sh, qn("a:srgbClr"))
        c.set("val", "0F1B3D")
        a = etree.SubElement(c, qn("a:alpha"))
        a.set("val", "10000")

    def rect(self, slide, r: Rect):
        kind = MSO_SHAPE.OVAL if r.shape == "ellipse" else (MSO_SHAPE.ROUNDED_RECTANGLE if r.radius > 0 else MSO_SHAPE.RECTANGLE)
        s = slide.shapes.add_shape(kind, Emu(int(r.x * PX)), Emu(int(r.y * PX)), Emu(int(r.w * PX)), Emu(int(r.h * PX)))
        if r.radius > 0 and r.shape != "ellipse":
            m = min(r.w, r.h)
            s.adjustments[0] = min(0.5, r.radius / m) if m else 0
        self._rgb_fill(s, r.fill, r.opacity)
        self._line(s, r.stroke, r.stroke_w, r.dash)
        if r.shadow:
            self._shadow(s)
        else:
            self._no_shadow(s)
        if r.name:
            s.name = r.name
        s.text_frame.text = ""
        return s

    def line(self, slide, l: Line):
        s = slide.shapes.add_connector(1, Emu(int(l.x1 * PX)), Emu(int(l.y1 * PX)), Emu(int(l.x2 * PX)), Emu(int(l.y2 * PX)))
        s.line.color.rgb = hexrgb(l.color)
        s.line.width = Emu(int(l.w * PX))
        ln = s.line._get_or_add_ln()
        if l.dash:
            pd = etree.SubElement(ln, qn("a:prstDash"))
            pd.set("val", "dash")
        if l.arrow:
            te = etree.SubElement(ln, qn("a:tailEnd"))
            te.set("type", "triangle")
            te.set("w", "med")
            te.set("len", "med")
        self._no_shadow(s)   # 커넥터도 테마 기본 그림자를 물려받으므로 명시적으로 끈다
        return s

    def image(self, slide, im: Image):
        p = im.path if os.path.isabs(im.path) else os.path.join(ASSETS, im.path)
        s = slide.shapes.add_picture(p, Emu(int(im.x * PX)), Emu(int(im.y * PX)), Emu(int(im.w * PX)), Emu(int(im.h * PX)))
        if im.name:
            s.name = im.name
        return s

    def icon(self, slide, ic: Icon):
        fn = f"icons/{ic.name}-{ic.color.lstrip('#').lower()}.png"
        return self.image(slide, Image(fn, ic.x, ic.y, ic.size, ic.size, name=f"icon:{ic.name}"))

    def text(self, slide, t: Text):
        tb = slide.shapes.add_textbox(Emu(int(t.x * PX)), Emu(int(t.y * PX)), Emu(int(t.w * PX)), Emu(int(t.h * PX)))
        if t.name:
            tb.name = t.name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}[t.valign]
        # autofit 끄기
        bodyPr = tf._txBody.find(qn("a:bodyPr"))
        for ch in list(bodyPr):
            bodyPr.remove(ch)
        etree.SubElement(bodyPr, qn("a:noAutofit"))
        ps = _paragraphs(t)
        for i, para in enumerate(ps):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}[t.align]
            p.line_spacing = t.line
            if t.para_gap and i > 0:
                p.space_before = Pt(t.para_gap * 0.75)
            if not para:
                r = p.add_run()
                r.text = ""
                self._style_run(r, t, Run(""))
                continue
            for run in para:
                r = p.add_run()
                r.text = run.text
                self._style_run(r, t, run)
        return tb

    def _style_run(self, r, t: Text, run: Run):
        USED_CHARS.update(run.text)
        size = run.size or t.size
        r.font.size = Pt(size * 0.75)  # px → pt
        weight = run.weight or t.weight
        r.font.bold = bool(run.bold or weight >= 600)
        r.font.color.rgb = hexrgb(run.color or t.color)
        r.font.name = FONT
        rPr = r._r.get_or_add_rPr()
        for tag in ("a:latin", "a:ea", "a:cs"):
            el = rPr.find(qn(tag))
            if el is None:
                el = etree.SubElement(rPr, qn(tag))
            el.set("typeface", FONT)

    def slide(self, sd: Slide):
        s = self.prs.slides.add_slide(self.blank)
        # 배경
        bg = s.background.fill
        bg.solid()
        bg.fore_color.rgb = hexrgb(sd.bg)
        for it in sd.items:
            if isinstance(it, Rect):
                self.rect(s, it)
            elif isinstance(it, Text):
                self.text(s, it)
            elif isinstance(it, Line):
                self.line(s, it)
            elif isinstance(it, Image):
                self.image(s, it)
            elif isinstance(it, Icon):
                self.icon(s, it)
        if sd.notes:
            s.notes_slide.notes_text_frame.text = sd.notes
        return s

    def save(self, path):
        self.prs.save(path)
        _set_theme_font(path, FONT)
        # 구글 슬라이드는 Noto Sans KR 을 자체 보유하므로 심을 필요가 없다.
        # 파워포인트로도 열어야 하면  EMBED_FONT=1 python3 build.py
        if os.environ.get("EMBED_FONT") == "1":
            try:
                import embed_font
                embed_font.embed(path, FONT, USED_CHARS)
            except Exception as e:
                print(f"[경고] 글꼴 embed 생략: {e}")


def _set_theme_font(path, font):
    """저장된 pptx 안의 테마 기본 글꼴(제목·본문)을 바꾼다."""
    import re
    import shutil
    import zipfile

    src = zipfile.ZipFile(path, "r")
    items = src.infolist()
    data = {i.filename: src.read(i.filename) for i in items}
    src.close()

    for name in list(data):
        if not re.match(r"ppt/theme/theme\d+\.xml$", name):
            continue
        xml = data[name].decode("utf-8")

        def fix(block):
            block = re.sub(r'<a:latin typeface="[^"]*"', f'<a:latin typeface="{font}"', block)
            block = re.sub(r'<a:ea typeface="[^"]*"', f'<a:ea typeface="{font}"', block)
            block = re.sub(r'<a:cs typeface="[^"]*"', f'<a:cs typeface="{font}"', block)
            return block

        for tag in ("majorFont", "minorFont"):
            m = re.search(rf"<a:{tag}>.*?</a:{tag}>", xml, re.S)
            if m:
                xml = xml[:m.start()] + fix(m.group(0)) + xml[m.end():]
        data[name] = xml.encode("utf-8")

    tmp = path + ".tmp"
    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as out:
        for i in items:
            out.writestr(i, data[i.filename])
    shutil.move(tmp, path)


# ---------------------------------------------------------------- html emit
class HtmlWriter:
    def __init__(self, icons_svg: dict):
        self.icons = icons_svg
        self.slides_html = []

    @staticmethod
    def _sty(**kw):
        return ";".join(f"{k.replace('_', '-')}:{v}" for k, v in kw.items() if v is not None)

    def rect(self, r: Rect):
        st = dict(position="absolute", left=f"{r.x}px", top=f"{r.y}px", width=f"{r.w}px", height=f"{r.h}px",
                  background=r.fill or "transparent", border_radius=("50%" if r.shape == "ellipse" else f"{r.radius}px"),
                  box_sizing="border-box", opacity=r.opacity if r.opacity < 1 else None)
        if r.stroke:
            st["border"] = f"{r.stroke_w}px {'dashed' if r.dash else 'solid'} {r.stroke}"
        if r.shadow:
            st["box_shadow"] = "0 4px 14px rgba(15,27,61,.10)"
        return f'<div style="{self._sty(**st)}"></div>'

    def line(self, l: Line):
        import math
        dx, dy = l.x2 - l.x1, l.y2 - l.y1
        length = math.hypot(dx, dy)
        ang = math.degrees(math.atan2(dy, dx))
        head = ""
        if l.arrow:
            head = f'<div style="position:absolute;right:-1px;top:50%;width:0;height:0;border-left:{l.w*4+4}px solid {l.color};border-top:{l.w*2+3}px solid transparent;border-bottom:{l.w*2+3}px solid transparent;transform:translateY(-50%)"></div>'
        bstyle = "dashed" if l.dash else "solid"
        return (f'<div style="position:absolute;left:{l.x1}px;top:{l.y1}px;width:{length}px;height:0;'
                f'border-top:{l.w}px {bstyle} {l.color};transform-origin:0 0;transform:rotate({ang}deg) translateY(-{l.w/2}px)">{head}</div>')

    def image(self, im: Image):
        p = im.path if os.path.isabs(im.path) else os.path.join(ASSETS, im.path)
        return f'<img src="file://{p}" style="position:absolute;left:{im.x}px;top:{im.y}px;width:{im.w}px;height:{im.h}px">'

    def icon(self, ic: Icon):
        svg = self.icons[ic.name]
        return (f'<div style="position:absolute;left:{ic.x}px;top:{ic.y}px;width:{ic.size}px;height:{ic.size}px;color:{ic.color}">'
                f'{svg}</div>')

    def text(self, t: Text):
        ps = _paragraphs(t)
        jc = {"top": "flex-start", "middle": "center", "bottom": "flex-end"}[t.valign]
        out = [f'<div style="position:absolute;left:{t.x}px;top:{t.y}px;width:{t.w}px;height:{t.h}px;display:flex;flex-direction:column;justify-content:{jc};'
               f'font-size:{t.size}px;color:{t.color};font-weight:{t.weight};text-align:{t.align};line-height:{t.line};letter-spacing:{t.letter}px;word-break:keep-all;overflow-wrap:break-word">']
        for i, para in enumerate(ps):
            mt = f"margin-top:{t.para_gap}px;" if (i > 0 and t.para_gap) else ""
            out.append(f'<div style="{mt}min-height:{t.size * t.line}px">')
            for run in para:
                st = []
                if run.bold or (run.weight and run.weight >= 600):
                    st.append("font-weight:700")
                elif run.weight:
                    st.append(f"font-weight:{run.weight}")
                if run.color:
                    st.append(f"color:{run.color}")
                if run.size:
                    st.append(f"font-size:{run.size}px")
                out.append(f'<span style="{";".join(st)}">{_html.escape(run.text)}</span>')
            out.append("</div>")
        out.append("</div>")
        return "".join(out)

    def slide(self, sd: Slide, idx: int):
        parts = [f'<section class="slide" id="s{idx}" style="background:{sd.bg}">']
        for it in sd.items:
            if isinstance(it, Rect):
                parts.append(self.rect(it))
            elif isinstance(it, Text):
                parts.append(self.text(it))
            elif isinstance(it, Line):
                parts.append(self.line(it))
            elif isinstance(it, Image):
                parts.append(self.image(it))
            elif isinstance(it, Icon):
                parts.append(self.icon(it))
        parts.append("</section>")
        self.slides_html.append("".join(parts))

    def save(self, path):
        doc = f"""<!doctype html><html lang="ko"><head><meta charset="utf-8"><title>preview</title>
<link rel="preconnect" href="https://fonts.googleapis.com">
<link href="https://fonts.googleapis.com/css2?family=Noto+Sans+KR:wght@400;500;700;900&display=swap" rel="stylesheet">
<style>
body{{margin:0;background:#666;font-family:'Noto Sans KR',sans-serif;-webkit-font-smoothing:antialiased}}
.slide{{position:relative;width:{W}px;height:{H}px;overflow:hidden;margin:0 0 24px 0}}
.slide *{{box-sizing:border-box}}
svg{{width:100%;height:100%;display:block}}
</style></head><body>
{''.join(self.slides_html)}
</body></html>"""
        with open(path, "w", encoding="utf-8") as f:
            f.write(doc)


BODY_TOP, BODY_BOTTOM = 142, 672


def _span(it):
    if isinstance(it, Line):
        return min(it.y1, it.y2), max(it.y1, it.y2)
    if isinstance(it, Icon):
        return it.y, it.y + it.size
    return it.y, it.y + it.h


def _move(it, dy):
    if isinstance(it, Line):
        it.y1 += dy
        it.y2 += dy
    else:
        it.y += dy


def balance_body(sd: Slide, min_bottom: float = 30, max_shift: float = 26) -> float:
    """아래가 비면 본문 전체를 그만큼 내려 제목·설명과의 간격을 넓힌다."""
    # 푸터(구분선 y=676, 문구 y=684)는 본문에서 제외한다.
    body = [it for it in sd.items if BODY_TOP <= _span(it)[0] and _span(it)[1] < BODY_BOTTOM]
    if not body:
        return 0
    slack = BODY_BOTTOM - max(_span(it)[1] for it in body)
    dy = min(max(slack - min_bottom, 0), max_shift)
    if dy <= 0:
        return 0
    for it in body:
        _move(it, dy)
    return dy


def _strip_theme_effects(prs) -> int:
    """효과가 지정되지 않은 도형·그림·글상자에 빈 effectLst를 넣어
    파워포인트 테마의 기본 그림자를 물려받지 않게 한다.
    이미 효과가 있는 도형(shadow=True)은 건드리지 않는다."""
    n = 0
    for sl in prs.slides:
        for sh in sl.shapes:
            el = sh._element
            spPr = el.find(qn("p:spPr"))
            if spPr is None or spPr.find(qn("a:effectLst")) is not None:
                continue
            etree.SubElement(spPr, qn("a:effectLst"))
            n += 1
    return n


def build(slides: list[Slide], pptx_path: str, html_path: str, icons_svg: dict):
    pw = PptxWriter()
    hw = HtmlWriter(icons_svg)
    for i, sd in enumerate(slides):
        if sd.balance:
            balance_body(sd)
        pw.slide(sd)
        hw.slide(sd, i + 1)
    _strip_theme_effects(pw.prs)
    pw.save(pptx_path)
    hw.save(html_path)
